"""Import aircraft secondary-weapon presets and quantities."""
import json
import pathlib
import sys
import openpyxl

ROOT=pathlib.Path(__file__).resolve().parents[1]
DEFAULT_SOURCE=ROOT.parents[1]/'upload'/'war_thunder_aircraft_secondary_weapons(1).xlsx'

def rows(book,sheet):
    values=book[sheet].iter_rows(values_only=True);headers=next(values)
    return[dict(zip(headers,row)) for row in values if any(value is not None for value in row)]

def main():
    source=pathlib.Path(sys.argv[1]) if len(sys.argv)>1 else DEFAULT_SOURCE
    book=openpyxl.load_workbook(source,read_only=True,data_only=True)
    payload={'meta':{'title':'Aircraft Secondary Weapons','snapshot':'2026-09-11','sourceWorkbook':source.name},'aircraft':rows(book,'Aircraft'),'loadoutDetail':rows(book,'Loadout Detail'),'ordnance':rows(book,'Ordnance'),'fireBombSearch':rows(book,'Fire Bomb Search'),'readme':[list(row) for row in book['READ ME'].iter_rows(values_only=True) if any(value is not None for value in row)]}
    assert len(payload['aircraft'])==1556 and len(payload['loadoutDetail'])==14217 and len(payload['fireBombSearch'])==324
    destination=ROOT/'data'/'aircraft_loadouts.json';destination.write_text(json.dumps(payload,ensure_ascii=False,separators=(',',':'),default=str),encoding='utf-8')
    print(f'{destination.name}: {destination.stat().st_size:,} bytes')

if __name__=='__main__':main()
